from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models import KPIResult, KPIDefinition
from app.services.analytics import (
    executive_trend,
    local_authority_rows_for_month,
    service_rows_for_month,
)


# Brand / report palette
NAVY = "173042"
BLUE = "1F5B83"
MID_BLUE = "2D709E"
PALE_BLUE = "EAF2F8"
VERY_PALE_BLUE = "F5F9FC"
GREEN = "E2F0D9"
GREEN_DARK = "23734A"
AMBER = "FFF2CC"
AMBER_DARK = "936100"
RED = "F4CCCC"
RED_DARK = "A62F2F"
GREY = "E7E6E6"
LIGHT_GREY = "F3F5F7"
MID_GREY = "D9E0E5"
DARK_GREY = "66737C"
WHITE = "FFFFFF"
BLACK = "1F2933"

THIN_GREY = Side(style="thin", color="D9E1E6")
MEDIUM_BLUE = Side(style="medium", color=BLUE)

LOGO_PATH = Path(__file__).resolve().parents[1] / "static" / "blue_ribbon_logo.png"


def _solid(color):
    return PatternFill("solid", fgColor=color)


def _add_logo(ws, anchor="A1", width=44, height=44):
    if not LOGO_PATH.exists():
        return

    try:
        img = XLImage(str(LOGO_PATH))
        img.width = width
        img.height = height
        ws.add_image(img, anchor)
    except Exception:
        # Excel export should not fail because of a logo issue.
        pass


def _title(ws, title, subtitle=None, end_col=8):
    """Professional two-row report banner with Blue Ribbon branding."""
    end_col = max(end_col, 8)
    ws.sheet_view.showGridLines = False

    for row in (1, 2):
        for col in range(1, end_col + 1):
            ws.cell(row=row, column=col).fill = _solid(NAVY)

    ws.merge_cells(
        start_row=1,
        start_column=2,
        end_row=1,
        end_column=end_col,
    )
    ws.merge_cells(
        start_row=2,
        start_column=2,
        end_row=2,
        end_column=end_col,
    )

    ws["B1"] = title
    ws["B1"].font = Font(
        name="Segoe UI",
        size=18,
        bold=True,
        color=WHITE,
    )
    ws["B1"].alignment = Alignment(vertical="center")

    ws["B2"] = subtitle or "Blue Ribbon Health & Wellbeing"
    ws["B2"].font = Font(
        name="Segoe UI",
        size=9,
        color="C9D9E4",
    )
    ws["B2"].alignment = Alignment(vertical="center")

    ws.row_dimensions[1].height = 29
    ws.row_dimensions[2].height = 22

    _add_logo(ws, "A1", 42, 42)

    for col in range(1, end_col + 1):
        ws.cell(row=3, column=col).fill = _solid(MID_BLUE)

    ws.row_dimensions[3].height = 4


def _section(ws, row, title, start_col=1, end_col=8):
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = _solid(PALE_BLUE)
        cell.border = Border(bottom=MEDIUM_BLUE)

    cell = ws.cell(row=row, column=start_col)
    cell.value = title
    cell.font = Font(
        name="Segoe UI",
        size=11,
        bold=True,
        color=NAVY,
    )
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 22


def _header(row_cells):
    for cell in row_cells:
        cell.font = Font(
            name="Segoe UI",
            size=9,
            bold=True,
            color=WHITE,
        )
        cell.fill = _solid(BLUE)
        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )
        cell.border = Border(
            bottom=Side(style="thin", color=WHITE)
        )


def _rag_fill(cell):
    value = str(cell.value or "").strip().lower()

    if value == "green":
        cell.fill = _solid(GREEN)
        cell.font = Font(
            name="Segoe UI",
            bold=True,
            color=GREEN_DARK,
        )

    elif value == "amber":
        cell.fill = _solid(AMBER)
        cell.font = Font(
            name="Segoe UI",
            bold=True,
            color=AMBER_DARK,
        )

    elif value == "red":
        cell.fill = _solid(RED)
        cell.font = Font(
            name="Segoe UI",
            bold=True,
            color=RED_DARK,
        )

    else:
        cell.fill = _solid(GREY)
        cell.font = Font(
            name="Segoe UI",
            bold=True,
            color=DARK_GREY,
        )


def _add_rag_conditional_formatting(ws, cell_range):
    first_cell = cell_range.split(":")[0]

    for text, fill, font_color in (
        ("Green", GREEN, GREEN_DARK),
        ("Amber", AMBER, AMBER_DARK),
        ("Red", RED, RED_DARK),
        ("Unscored", GREY, DARK_GREY),
    ):
        ws.conditional_formatting.add(
            cell_range,
            FormulaRule(
                formula=[f'{first_cell}="{text}"'],
                fill=_solid(fill),
                font=Font(
                    color=font_color,
                    bold=True,
                ),
            ),
        )


def _zebra(ws, start_row, end_row, start_col=1, end_col=None):
    if end_row < start_row:
        return

    end_col = end_col or ws.max_column

    for row in range(start_row, end_row + 1):
        if (row - start_row) % 2:
            for col in range(start_col, end_col + 1):
                cell = ws.cell(row=row, column=col)

                if cell.fill.fill_type is None:
                    cell.fill = _solid(VERY_PALE_BLUE)


def _autosize(ws, max_width=35, min_width=10):
    for idx in range(1, ws.max_column + 1):
        letter = get_column_letter(idx)
        width = 0

        for cell in ws[letter]:
            if cell.value is not None:
                width = max(width, len(str(cell.value)))

        ws.column_dimensions[letter].width = min(
            max(width + 2, min_width),
            max_width,
        )


def _rag_formula(score_ref):
    return (
        f'=IF({score_ref}="",'
        f'"Unscored",'
        f'IF({score_ref}>=2.5,"Green",'
        f'IF({score_ref}>=1.75,"Amber","Red")))'
    )


def _result_display(result):
    if result is None:
        return ""

    if result.value_numeric is not None:
        return result.value_numeric

    return result.value_text or ""


def _style_card(
    ws,
    label_cell,
    value_cell,
    fill=WHITE,
    accent=BLUE,
    value_format=None,
):
    label_cell.font = Font(
        name="Segoe UI",
        size=9,
        bold=True,
        color=DARK_GREY,
    )
    label_cell.fill = _solid(fill)
    label_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )
    label_cell.border = Border(
        top=MEDIUM_BLUE,
        left=THIN_GREY,
        right=THIN_GREY,
    )

    value_cell.fill = _solid(fill)
    value_cell.font = Font(
        name="Segoe UI",
        size=22,
        bold=True,
        color=accent,
    )
    value_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )
    value_cell.border = Border(
        left=THIN_GREY,
        right=THIN_GREY,
        bottom=THIN_GREY,
    )

    if value_format:
        value_cell.number_format = value_format


def _link_to_sheet(cell, sheet_title, display=None):
    cell.hyperlink = f"#'{sheet_title}'!A1"

    if display is not None:
        cell.value = display

    cell.style = "Hyperlink"
    cell.font = Font(
        name="Segoe UI",
        size=9,
        color=BLUE,
        underline="single",
    )


def _finish_sheet(
    ws,
    landscape=True,
    fit_to_width=1,
    print_title_rows=None,
):
    ws.sheet_view.showGridLines = False

    ws.page_setup.orientation = (
        "landscape" if landscape else "portrait"
    )
    ws.page_setup.fitToWidth = fit_to_width
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws.print_options.horizontalCentered = True

    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5

    ws.oddFooter.left.text = "Blue Ribbon Balanced Scorecard"
    ws.oddFooter.left.size = 8
    ws.oddFooter.left.color = DARK_GREY

    ws.oddFooter.center.text = "CONFIDENTIAL"
    ws.oddFooter.center.size = 8
    ws.oddFooter.center.color = DARK_GREY

    ws.oddFooter.right.text = "Page &P of &N"
    ws.oddFooter.right.size = 8
    ws.oddFooter.right.color = DARK_GREY

    if print_title_rows:
        ws.print_title_rows = print_title_rows


def build_board_pack(reporting_month):
    wb = Workbook()

    wb.properties.title = (
        f"Blue Ribbon Balanced Scorecard - "
        f"{reporting_month:%B %Y}"
    )
    wb.properties.subject = "Monthly Board Performance Dashboard"
    wb.properties.creator = "Blue Ribbon Health & Wellbeing"

    ws = wb.active
    ws.title = "Executive Summary"
    ws.sheet_properties.tabColor = NAVY

    service_rows = service_rows_for_month(reporting_month)
    local_authority_rows = local_authority_rows_for_month(
        reporting_month
    )

    kpis = (
        KPIDefinition.query
        .filter_by(
            active=True,
            group_only=False,
        )
        .order_by(
            KPIDefinition.domain,
            KPIDefinition.name,
        )
        .all()
    )

    group_results = (
        KPIResult.query
        .filter_by(
            reporting_month=reporting_month,
            scope="group",
        )
        .join(KPIDefinition)
        .order_by(
            KPIDefinition.domain,
            KPIDefinition.name,
        )
        .all()
    )

    # ============================================================
    # ALL SERVICES
    # ============================================================

    all_ws = wb.create_sheet("All Services")
    all_ws.sheet_properties.tabColor = MID_BLUE

    headers = [
        "Service",
        "Regional Manager",
        "Registered Manager",
        "Local Authority",
        "Overall Score",
        "Overall RAG",
    ] + [
        kpi.name for kpi in kpis
    ] + [
        "Red KPIs",
        "Amber KPIs",
        "Sort Key",
        "Service Type",
    ]

    _title(
        all_ws,
        "Blue Ribbon Balanced Scorecard - Service Matrix",
        reporting_month.strftime(
            "Reporting month: %B %Y"
        ),
        end_col=min(
            max(len(headers), 8),
            18,
        ),
    )

    for col, value in enumerate(headers, 1):
        all_ws.cell(
            row=4,
            column=col,
            value=value,
        )

    _header(all_ws[4])
    all_ws.row_dimensions[4].height = 34

    first_kpi_col = 7
    last_kpi_col = first_kpi_col + len(kpis) - 1

    all_ws.cell(
        row=5,
        column=1,
        value="KPI weights",
    )

    for offset, kpi in enumerate(kpis):
        all_ws.cell(
            row=5,
            column=first_kpi_col + offset,
            value=kpi.weight or 1.0,
        )

    all_ws.row_dimensions[5].hidden = True
    all_ws.freeze_panes = "G6"

    service_sheet_titles = {}
    used_titles = set(wb.sheetnames)

    for row in service_rows:
        raw_title = row["service"].name[:31]
        title = raw_title
        suffix = 2

        while title in used_titles:
            suffix_text = f" {suffix}"
            title = (
                raw_title[
                    :31 - len(suffix_text)
                ]
                + suffix_text
            )
            suffix += 1

        used_titles.add(title)
        service_sheet_titles[
            row["service"].id
        ] = title

    for excel_row, row in enumerate(service_rows, 6):
        service = row["service"]

        all_ws.cell(
            row=excel_row,
            column=1,
            value=service.name,
        )

        _link_to_sheet(
            all_ws.cell(
                row=excel_row,
                column=1,
            ),
            service_sheet_titles[service.id],
            service.name,
        )

        all_ws.cell(
            row=excel_row,
            column=2,
            value=service.regional_manager or "Unassigned",
        )

        all_ws.cell(
            row=excel_row,
            column=3,
            value=service.registered_manager or "Unassigned",
        )

        all_ws.cell(
            row=excel_row,
            column=4,
            value=service.local_authority or "Unassigned",
        )

        result_by_kpi = {
            result.kpi_id: result
            for result in row["results"]
        }

        for offset, kpi in enumerate(kpis):
            result = result_by_kpi.get(kpi.id)

            target_cell = all_ws.cell(
                row=excel_row,
                column=first_kpi_col + offset,
            )

            if result:
                target_cell.value = result.rag
                _rag_fill(target_cell)
            else:
                target_cell.value = "Unscored"
                _rag_fill(target_cell)

        if kpis:
            score_parts = []

            for offset, _kpi in enumerate(kpis):
                col_letter = get_column_letter(
                    first_kpi_col + offset
                )
                weight_cell = f"{col_letter}$5"
                rag_cell = f"{col_letter}{excel_row}"

                score_parts.append(
                    f'IF({rag_cell}="Green",3*{weight_cell},'
                    f'IF({rag_cell}="Amber",2*{weight_cell},'
                    f'IF({rag_cell}="Red",1*{weight_cell},0)))'
                )

            weight_parts = []

            for offset, _kpi in enumerate(kpis):
                col_letter = get_column_letter(
                    first_kpi_col + offset
                )
                weight_cell = f"{col_letter}$5"
                rag_cell = f"{col_letter}{excel_row}"

                weight_parts.append(
                    f'IF(OR({rag_cell}="Green",'
                    f'{rag_cell}="Amber",'
                    f'{rag_cell}="Red"),'
                    f'{weight_cell},0)'
                )

            all_ws.cell(
                row=excel_row,
                column=5,
                value=(
                    f'=IFERROR(('
                    f'{"+".join(score_parts)}'
                    f')/('
                    f'{"+".join(weight_parts)}'
                    f'),"")'
                ),
            )

            all_ws.cell(
                row=excel_row,
                column=5,
            ).number_format = "0.00"

        all_ws.cell(
            row=excel_row,
            column=6,
            value=_rag_formula(
                f"E{excel_row}"
            ),
        )

        _rag_fill(
            all_ws.cell(
                row=excel_row,
                column=6,
            )
        )

        red_col = last_kpi_col + 1
        amber_col = last_kpi_col + 2
        sort_col = last_kpi_col + 3
        type_col = last_kpi_col + 4

        first_letter = get_column_letter(
            first_kpi_col
        )
        last_letter = get_column_letter(
            last_kpi_col
        )

        if kpis:
            all_ws.cell(
                row=excel_row,
                column=red_col,
                value=(
                    f'=COUNTIF('
                    f'{first_letter}{excel_row}:'
                    f'{last_letter}{excel_row},'
                    f'"Red")'
                ),
            )

            all_ws.cell(
                row=excel_row,
                column=amber_col,
                value=(
                    f'=COUNTIF('
                    f'{first_letter}{excel_row}:'
                    f'{last_letter}{excel_row},'
                    f'"Amber")'
                ),
            )

        all_ws.cell(
            row=excel_row,
            column=sort_col,
            value=(
                f'=IF(F{excel_row}="Red",1,'
                f'IF(F{excel_row}="Amber",2,'
                f'IF(F{excel_row}="Green",3,4)))'
            ),
        )

        all_ws.cell(
            row=excel_row,
            column=type_col,
            value=service.service_type or "Not set",
        )

    if service_rows:
        _add_rag_conditional_formatting(
            all_ws,
            f"F6:F{5 + len(service_rows)}",
        )

        if kpis:
            _add_rag_conditional_formatting(
                all_ws,
                (
                    f"{get_column_letter(first_kpi_col)}6:"
                    f"{get_column_letter(last_kpi_col)}"
                    f"{5 + len(service_rows)}"
                ),
            )

        _zebra(
            all_ws,
            6,
            5 + len(service_rows),
            1,
            len(headers),
        )

    _autosize(all_ws, 32)

    all_ws.column_dimensions["A"].width = 26
    all_ws.column_dimensions["B"].width = 24
    all_ws.column_dimensions["C"].width = 24
    all_ws.column_dimensions["D"].width = 22
    all_ws.column_dimensions["E"].width = 14
    all_ws.column_dimensions["F"].width = 13

    all_ws.auto_filter.ref = (
        f"A4:{get_column_letter(len(headers))}"
        f"{max(4, 5 + len(service_rows))}"
    )

    _finish_sheet(
        all_ws,
        landscape=True,
        fit_to_width=1,
        print_title_rows="1:4",
    )

    # ============================================================
    # EXECUTIVE SUMMARY
    # ============================================================

    _title(
        ws,
        "Blue Ribbon Balanced Scorecard",
        reporting_month.strftime(
            "Executive Summary | %B %Y"
        ),
        end_col=12,
    )

    ws.column_dimensions["A"].width = 3

    for col in range(2, 13):
        ws.column_dimensions[
            get_column_letter(col)
        ].width = 14

    total_services = len(service_rows)

    green_services = sum(
        1
        for row in service_rows
        if str(row.get("rag") or "").lower() == "green"
    )
    amber_services = sum(
        1
        for row in service_rows
        if str(row.get("rag") or "").lower() == "amber"
    )
    red_services = sum(
        1
        for row in service_rows
        if str(row.get("rag") or "").lower() == "red"
    )
    unscored_services = (
        total_services
        - green_services
        - amber_services
        - red_services
    )

    scored_rows = [
        row
        for row in service_rows
        if row.get("score") is not None
    ]

    overall_score = None

    if scored_rows:
        overall_score = sum(
            row["score"]
            for row in scored_rows
        ) / len(scored_rows)

    ws["B5"] = "OVERALL SCORE"
    ws["B6"] = overall_score or ""

    _style_card(
        ws,
        ws["B5"],
        ws["B6"],
        fill=VERY_PALE_BLUE,
        accent=NAVY,
        value_format="0.00",
    )

    ws["D5"] = "SERVICES"
    ws["D6"] = total_services

    _style_card(
        ws,
        ws["D5"],
        ws["D6"],
        fill=WHITE,
        accent=BLUE,
        value_format="0",
    )

    ws["F5"] = "GREEN"
    ws["F6"] = green_services

    _style_card(
        ws,
        ws["F5"],
        ws["F6"],
        fill=GREEN,
        accent=GREEN_DARK,
        value_format="0",
    )

    ws["H5"] = "AMBER"
    ws["H6"] = amber_services

    _style_card(
        ws,
        ws["H5"],
        ws["H6"],
        fill=AMBER,
        accent=AMBER_DARK,
        value_format="0",
    )

    ws["J5"] = "RED"
    ws["J6"] = red_services

    _style_card(
        ws,
        ws["J5"],
        ws["J6"],
        fill=RED,
        accent=RED_DARK,
        value_format="0",
    )

    ws["L5"] = "UNSCORED"
    ws["L6"] = unscored_services

    _style_card(
        ws,
        ws["L5"],
        ws["L6"],
        fill=GREY,
        accent=DARK_GREY,
        value_format="0",
    )

    ws.row_dimensions[5].height = 22
    ws.row_dimensions[6].height = 34

    # RAG distribution data
    rag_data_start = 9

    ws.cell(
        row=rag_data_start,
        column=2,
        value="RAG",
    )
    ws.cell(
        row=rag_data_start,
        column=3,
        value="Services",
    )

    rag_values = [
        ("Green", green_services),
        ("Amber", amber_services),
        ("Red", red_services),
        ("Unscored", unscored_services),
    ]

    for offset, (name, value) in enumerate(
        rag_values,
        rag_data_start + 1,
    ):
        ws.cell(
            row=offset,
            column=2,
            value=name,
        )
        ws.cell(
            row=offset,
            column=3,
            value=value,
        )

    pie = PieChart()
    pie.title = "Service RAG Distribution"

    labels = Reference(
        ws,
        min_col=2,
        min_row=rag_data_start + 1,
        max_row=rag_data_start + 4,
    )

    data = Reference(
        ws,
        min_col=3,
        min_row=rag_data_start,
        max_row=rag_data_start + 4,
    )

    pie.add_data(
        data,
        titles_from_data=True,
    )
    pie.set_categories(labels)
    pie.height = 7.5
    pie.width = 10
    pie.legend.position = "r"

    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True

    ws.add_chart(
        pie,
        "B15",
    )

    # Top / bottom services
    ranked = [
        row
        for row in service_rows
        if row.get("score") is not None
    ]

    ranked = sorted(
        ranked,
        key=lambda row: row["score"],
    )

    bottom_five = ranked[:5]
    top_five = list(
        reversed(ranked[-5:])
    )

    _section(
        ws,
        9,
        "KEY SERVICE EXCEPTIONS",
        5,
        12,
    )

    ws["E10"] = "Lowest scoring services"
    ws["E10"].font = Font(
        bold=True,
        color=NAVY,
    )

    headers = [
        "Service",
        "Score",
        "RAG",
    ]

    for col, header in enumerate(headers, 5):
        ws.cell(
            row=11,
            column=col,
            value=header,
        )

    _header(
        ws[11][4:7]
    )

    for idx, row in enumerate(
        bottom_five,
        12,
    ):
        ws.cell(
            row=idx,
            column=5,
            value=row["service"].name,
        )
        ws.cell(
            row=idx,
            column=6,
            value=row["score"],
        )
        ws.cell(
            row=idx,
            column=6,
        ).number_format = "0.00"
        ws.cell(
            row=idx,
            column=7,
            value=row["rag"],
        )
        _rag_fill(
            ws.cell(
                row=idx,
                column=7,
            )
        )

    ws["I10"] = "Highest scoring services"
    ws["I10"].font = Font(
        bold=True,
        color=NAVY,
    )

    for col, header in enumerate(headers, 9):
        ws.cell(
            row=11,
            column=col,
            value=header,
        )

    _header(
        ws[11][8:11]
    )

    for idx, row in enumerate(
        top_five,
        12,
    ):
        ws.cell(
            row=idx,
            column=9,
            value=row["service"].name,
        )
        ws.cell(
            row=idx,
            column=10,
            value=row["score"],
        )
        ws.cell(
            row=idx,
            column=10,
        ).number_format = "0.00"
        ws.cell(
            row=idx,
            column=11,
            value=row["rag"],
        )
        _rag_fill(
            ws.cell(
                row=idx,
                column=11,
            )
        )

    trend = executive_trend(
        reporting_month,
    )

    if trend:
        trend_start = 31

        _section(
            ws,
            trend_start,
            "12-MONTH OVERALL SCORE TREND",
            2,
            12,
        )

        ws.cell(
            row=trend_start + 1,
            column=2,
            value="Month",
        )
        ws.cell(
            row=trend_start + 1,
            column=3,
            value="Score",
        )

        for idx, point in enumerate(
            trend,
            trend_start + 2,
        ):
            ws.cell(
                row=idx,
                column=2,
                value=point["month"],
            )
            ws.cell(
                row=idx,
                column=3,
                value=point["score"],
            )

        chart = LineChart()
        chart.title = "Overall Score Trend"
        chart.y_axis.title = "Score"
        chart.x_axis.title = "Month"
        chart.height = 8
        chart.width = 20

        data = Reference(
            ws,
            min_col=3,
            min_row=trend_start + 1,
            max_row=trend_start + 1 + len(trend),
        )

        cats = Reference(
            ws,
            min_col=2,
            min_row=trend_start + 2,
            max_row=trend_start + 1 + len(trend),
        )

        chart.add_data(
            data,
            titles_from_data=True,
        )
        chart.set_categories(cats)

        ws.add_chart(
            chart,
            f"E{trend_start + 1}",
        )

    _finish_sheet(
        ws,
        landscape=True,
        fit_to_width=1,
    )

    # ============================================================
    # LOCAL AUTHORITY SUMMARY
    # ============================================================

    la_ws = wb.create_sheet(
        "Local Authority Summary"
    )
    la_ws.sheet_properties.tabColor = "5B9BD5"

    _title(
        la_ws,
        "Blue Ribbon Balanced Scorecard - Local Authority Summary",
        reporting_month.strftime(
            "Reporting month: %B %Y"
        ),
        end_col=8,
    )

    la_headers = [
        "Local Authority",
        "Services",
        "Average Score",
        "Overall RAG",
        "Red Services",
        "Amber Services",
    ]

    for col, value in enumerate(
        la_headers,
        1,
    ):
        la_ws.cell(
            row=4,
            column=col,
            value=value,
        )

    _header(
        la_ws[4]
    )

    authority_names = sorted(
        {
            row["service"].local_authority or "Unassigned"
            for row in service_rows
        }
    )

    for idx, authority in enumerate(
        authority_names,
        5,
    ):
        la_ws.cell(
            row=idx,
            column=1,
            value=authority,
        )

        la_ws.cell(
            row=idx,
            column=2,
            value=(
                f'=COUNTIF('
                f"'All Services'!$D:$D,"
                f"A{idx})"
            ),
        )

        la_ws.cell(
            row=idx,
            column=3,
            value=(
                f'=IFERROR('
                f'AVERAGEIF('
                f"'All Services'!$D:$D,"
                f"A{idx},"
                f"'All Services'!$E:$E"
                f'),"")'
            ),
        )

        la_ws.cell(
            row=idx,
            column=3,
        ).number_format = "0.00"

        la_ws.cell(
            row=idx,
            column=4,
            value=_rag_formula(
                f"C{idx}"
            ),
        )

        la_ws.cell(
            row=idx,
            column=5,
            value=(
                f'=COUNTIFS('
                f"'All Services'!$D:$D,"
                f"A{idx},"
                f"'All Services'!$F:$F,"
                f'"Red")'
            ),
        )

        la_ws.cell(
            row=idx,
            column=6,
            value=(
                f'=COUNTIFS('
                f"'All Services'!$D:$D,"
                f"A{idx},"
                f"'All Services'!$F:$F,"
                f'"Amber")'
            ),
        )

    if authority_names:
        _add_rag_conditional_formatting(
            la_ws,
            f"D5:D{4 + len(authority_names)}",
        )

        _zebra(
            la_ws,
            5,
            4 + len(authority_names),
            1,
            6,
        )

        la_ws.conditional_formatting.add(
            f"C5:C{4 + len(authority_names)}",
            ColorScaleRule(
                start_type="num",
                start_value=1,
                start_color="F8696B",
                mid_type="num",
                mid_value=2,
                mid_color="FFEB84",
                end_type="num",
                end_value=3,
                end_color="63BE7B",
            ),
        )

    _autosize(
        la_ws,
        30,
    )

    la_ws.freeze_panes = "A5"

    la_ws.auto_filter.ref = (
        f"A4:F{max(4, 4 + len(authority_names))}"
    )

    _finish_sheet(
        la_ws,
        landscape=True,
        fit_to_width=1,
        print_title_rows="1:4",
    )

    # ============================================================
    # REGIONAL MANAGER SUMMARY
    # ============================================================

    rm_ws = wb.create_sheet(
        "Regional Manager Summary"
    )
    rm_ws.sheet_properties.tabColor = "70AD47"

    _title(
        rm_ws,
        "Blue Ribbon Balanced Scorecard - Regional Manager Summary",
        reporting_month.strftime(
            "Reporting month: %B %Y"
        ),
        end_col=8,
    )

    rm_headers = [
        "Regional Manager",
        "Services",
        "Average Score",
        "Overall RAG",
        "Red Services",
        "Amber Services",
    ]

    for col, value in enumerate(
        rm_headers,
        1,
    ):
        rm_ws.cell(
            row=4,
            column=col,
            value=value,
        )

    _header(
        rm_ws[4]
    )

    manager_names = sorted(
        {
            row["service"].regional_manager or "Unassigned"
            for row in service_rows
        }
    )

    for idx, manager in enumerate(
        manager_names,
        5,
    ):
        rm_ws.cell(
            row=idx,
            column=1,
            value=manager,
        )

        rm_ws.cell(
            row=idx,
            column=2,
            value=(
                f'=COUNTIF('
                f"'All Services'!$B:$B,"
                f"A{idx})"
            ),
        )

        rm_ws.cell(
            row=idx,
            column=3,
            value=(
                f'=IFERROR('
                f'AVERAGEIF('
                f"'All Services'!$B:$B,"
                f"A{idx},"
                f"'All Services'!$E:$E"
                f'),"")'
            ),
        )

        rm_ws.cell(
            row=idx,
            column=3,
        ).number_format = "0.00"

        rm_ws.cell(
            row=idx,
            column=4,
            value=_rag_formula(
                f"C{idx}"
            ),
        )

        rm_ws.cell(
            row=idx,
            column=5,
            value=(
                f'=COUNTIFS('
                f"'All Services'!$B:$B,"
                f"A{idx},"
                f"'All Services'!$F:$F,"
                f'"Red")'
            ),
        )

        rm_ws.cell(
            row=idx,
            column=6,
            value=(
                f'=COUNTIFS('
                f"'All Services'!$B:$B,"
                f"A{idx},"
                f"'All Services'!$F:$F,"
                f'"Amber")'
            ),
        )

    if manager_names:
        _add_rag_conditional_formatting(
            rm_ws,
            f"D5:D{4 + len(manager_names)}",
        )

        _zebra(
            rm_ws,
            5,
            4 + len(manager_names),
            1,
            6,
        )

        rm_ws.conditional_formatting.add(
            f"C5:C{4 + len(manager_names)}",
            ColorScaleRule(
                start_type="num",
                start_value=1,
                start_color="F8696B",
                mid_type="num",
                mid_value=2,
                mid_color="FFEB84",
                end_type="num",
                end_value=3,
                end_color="63BE7B",
            ),
        )

    _autosize(
        rm_ws,
        30,
    )

    rm_ws.freeze_panes = "A5"

    rm_ws.auto_filter.ref = (
        f"A4:F{max(4, 4 + len(manager_names))}"
    )

    _finish_sheet(
        rm_ws,
        landscape=True,
        fit_to_width=1,
        print_title_rows="1:4",
    )

    # ============================================================
    # REGISTERED MANAGER SUMMARY
    # ============================================================

    rgm_ws = wb.create_sheet(
        "Registered Manager Summary"
    )
    rgm_ws.sheet_properties.tabColor = "A5A5A5"

    _title(
        rgm_ws,
        "Blue Ribbon Balanced Scorecard - Registered Manager Summary",
        reporting_month.strftime(
            "Reporting month: %B %Y"
        ),
        end_col=8,
    )

    rgm_headers = [
        "Regional Manager",
        "Registered Manager",
        "Services",
        "Average Score",
        "Overall RAG",
        "Red Services",
        "Amber Services",
    ]

    for col, value in enumerate(
        rgm_headers,
        1,
    ):
        rgm_ws.cell(
            row=4,
            column=col,
            value=value,
        )

    _header(
        rgm_ws[4]
    )

    manager_pairs = sorted(
        {
            (
                row["service"].regional_manager or "Unassigned",
                row["service"].registered_manager or "Unassigned",
            )
            for row in service_rows
        }
    )

    for idx, (
        regional_manager,
        registered_manager,
    ) in enumerate(
        manager_pairs,
        5,
    ):
        rgm_ws.cell(
            row=idx,
            column=1,
            value=regional_manager,
        )

        rgm_ws.cell(
            row=idx,
            column=2,
            value=registered_manager,
        )

        rgm_ws.cell(
            row=idx,
            column=3,
            value=(
                f"=COUNTIFS("
                f"'All Services'!$B:$B,A{idx},"
                f"'All Services'!$C:$C,B{idx})"
            ),
        )

        rgm_ws.cell(
            row=idx,
            column=4,
            value=(
                f'=IFERROR('
                f'AVERAGEIFS('
                f"'All Services'!$E:$E,"
                f"'All Services'!$B:$B,A{idx},"
                f"'All Services'!$C:$C,B{idx}"
                f'),"")'
            ),
        )

        rgm_ws.cell(
            row=idx,
            column=4,
        ).number_format = "0.00"

        rgm_ws.cell(
            row=idx,
            column=5,
            value=_rag_formula(
                f"D{idx}"
            ),
        )

        rgm_ws.cell(
            row=idx,
            column=6,
            value=(
                f'=COUNTIFS('
                f"'All Services'!$B:$B,A{idx},"
                f"'All Services'!$C:$C,B{idx},"
                f"'All Services'!$F:$F,\"Red\")"
            ),
        )

        rgm_ws.cell(
            row=idx,
            column=7,
            value=(
                f'=COUNTIFS('
                f"'All Services'!$B:$B,A{idx},"
                f"'All Services'!$C:$C,B{idx},"
                f"'All Services'!$F:$F,\"Amber\")"
            ),
        )

    if manager_pairs:
        _add_rag_conditional_formatting(
            rgm_ws,
            f"E5:E{4 + len(manager_pairs)}",
        )

        _zebra(
            rgm_ws,
            5,
            4 + len(manager_pairs),
            1,
            7,
        )

        rgm_ws.conditional_formatting.add(
            f"D5:D{4 + len(manager_pairs)}",
            ColorScaleRule(
                start_type="num",
                start_value=1,
                start_color="F8696B",
                mid_type="num",
                mid_value=2,
                mid_color="FFEB84",
                end_type="num",
                end_value=3,
                end_color="63BE7B",
            ),
        )

    _autosize(
        rgm_ws,
        30,
    )

    rgm_ws.freeze_panes = "A5"

    rgm_ws.auto_filter.ref = (
        f"A4:G{max(4, 4 + len(manager_pairs))}"
    )

    _finish_sheet(
        rgm_ws,
        landscape=True,
        fit_to_width=1,
        print_title_rows="1:4",
    )

    # ============================================================
    # GROUP KPIs
    # ============================================================

    grp_ws = wb.create_sheet(
        "Group KPIs"
    )
    grp_ws.sheet_properties.tabColor = "FFC000"

    _title(
        grp_ws,
        "Blue Ribbon Balanced Scorecard - Group Measures",
        reporting_month.strftime(
            "Reporting month: %B %Y"
        ),
        end_col=8,
    )

    grp_headers = [
        "Domain",
        "KPI",
        "Value",
        "RAG",
        "Commentary",
        "Source",
    ]

    for col, value in enumerate(
        grp_headers,
        1,
    ):
        grp_ws.cell(
            row=4,
            column=col,
            value=value,
        )

    _header(
        grp_ws[4]
    )

    for idx, result in enumerate(
        group_results,
        5,
    ):
        grp_ws.cell(
            row=idx,
            column=1,
            value=result.kpi.domain,
        )

        grp_ws.cell(
            row=idx,
            column=2,
            value=result.kpi.name,
        )

        grp_ws.cell(
            row=idx,
            column=3,
            value=_result_display(result),
        )

        if result.value_numeric is not None:
            if result.kpi.unit == "%":
                grp_ws.cell(
                    row=idx,
                    column=3,
                ).number_format = (
                    '0.0"%";-0.0"%";0"%"'
                )
            else:
                grp_ws.cell(
                    row=idx,
                    column=3,
                ).number_format = (
                    '0.0;-0.0;0'
                )

        grp_ws.cell(
            row=idx,
            column=4,
            value=result.rag,
        )

        _rag_fill(
            grp_ws.cell(
                row=idx,
                column=4,
            )
        )

        grp_ws.cell(
            row=idx,
            column=5,
            value=result.commentary or "",
        )

        grp_ws.cell(
            row=idx,
            column=6,
            value=result.source or "",
        )

        grp_ws.cell(
            row=idx,
            column=5,
        ).alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

        for col in range(1, 7):
            grp_ws.cell(
                row=idx,
                column=col,
            ).border = Border(
                bottom=THIN_GREY
            )

    if group_results:
        _zebra(
            grp_ws,
            5,
            4 + len(group_results),
            1,
            6,
        )

    _autosize(
        grp_ws,
        45,
    )

    grp_ws.column_dimensions["E"].width = 48
    grp_ws.freeze_panes = "A5"

    grp_ws.auto_filter.ref = (
        f"A4:F{max(4, 4 + len(group_results))}"
    )

    _finish_sheet(
        grp_ws,
        landscape=True,
        fit_to_width=1,
        print_title_rows="1:4",
    )

    # ============================================================
    # SERVICE DETAIL SHEETS
    # ============================================================

    for row in service_rows:
        service = row["service"]
        title = service_sheet_titles[
            service.id
        ]

        svc_ws = wb.create_sheet(
            title
        )

        svc_ws.sheet_properties.tabColor = "D9EAD3"

        _title(
            svc_ws,
            f"Service Scorecard - {service.name}",
            reporting_month.strftime(
                "Reporting month: %B %Y"
            ),
            end_col=10,
        )

        # IMPORTANT:
        # Row 2 is merged B2:J2 by _title(), so the link must NOT go in J2.
        svc_ws["J4"] = "Back to Executive Summary"

        _link_to_sheet(
            svc_ws["J4"],
            "Executive Summary",
            "Back to Executive Summary",
        )

        svc_ws["J4"].alignment = Alignment(
            horizontal="right",
            vertical="center",
        )

        # Service profile
        profile = [
            (
                "Regional Manager",
                service.regional_manager or "Not set",
            ),
            (
                "Registered Manager",
                service.registered_manager or "Not set",
            ),
            (
                "Local Authority",
                service.local_authority or "Unassigned",
            ),
            (
                "Service Type",
                service.service_type or "Not set",
            ),
        ]

        col_pairs = [
            (1, 2),
            (3, 4),
            (5, 6),
            (7, 8),
        ]

        for (
            label,
            value,
        ), (
            c1,
            c2,
        ) in zip(
            profile,
            col_pairs,
        ):
            svc_ws.cell(
                row=4,
                column=c1,
                value=label,
            )

            svc_ws.cell(
                row=4,
                column=c1,
            ).font = Font(
                name="Segoe UI",
                size=8,
                bold=True,
                color=DARK_GREY,
            )

            svc_ws.cell(
                row=4,
                column=c2,
                value=value,
            )

            svc_ws.cell(
                row=4,
                column=c2,
            ).font = Font(
                name="Segoe UI",
                size=9,
                bold=True,
                color=NAVY,
            )

            svc_ws.cell(
                row=4,
                column=c1,
            ).fill = _solid(
                PALE_BLUE
            )

            svc_ws.cell(
                row=4,
                column=c2,
            ).fill = _solid(
                PALE_BLUE
            )

        _section(
            svc_ws,
            6,
            "KPI PERFORMANCE",
            1,
            8,
        )

        service_headers = [
            "Domain",
            "KPI",
            "Value",
            "RAG",
            "Commentary",
            "Source",
        ]

        for col, value in enumerate(
            service_headers,
            1,
        ):
            svc_ws.cell(
                row=7,
                column=col,
                value=value,
            )

        _header(
            svc_ws[7][0:6]
        )

        ordered = sorted(
            row["results"],
            key=lambda r: (
                r.kpi.domain,
                r.kpi.name,
            ),
        )

        for idx, result in enumerate(
            ordered,
            8,
        ):
            svc_ws.cell(
                row=idx,
                column=1,
                value=result.kpi.domain,
            )

            svc_ws.cell(
                row=idx,
                column=2,
                value=result.kpi.name,
            )

            svc_ws.cell(
                row=idx,
                column=3,
                value=_result_display(result),
            )

            if result.value_numeric is not None:
                if result.kpi.unit == "%":
                    svc_ws.cell(
                        row=idx,
                        column=3,
                    ).number_format = (
                        '0.0"%";-0.0"%";0"%"'
                    )
                else:
                    svc_ws.cell(
                        row=idx,
                        column=3,
                    ).number_format = (
                        '0.0;-0.0;0'
                    )

            svc_ws.cell(
                row=idx,
                column=4,
                value=result.rag,
            )

            _rag_fill(
                svc_ws.cell(
                    row=idx,
                    column=4,
                )
            )

            svc_ws.cell(
                row=idx,
                column=5,
                value=result.commentary or "",
            )

            svc_ws.cell(
                row=idx,
                column=6,
                value=result.source or "",
            )

            svc_ws.cell(
                row=idx,
                column=5,
            ).alignment = Alignment(
                wrap_text=True,
                vertical="top",
            )

            for col in range(1, 7):
                svc_ws.cell(
                    row=idx,
                    column=col,
                ).border = Border(
                    bottom=THIN_GREY
                )

                svc_ws.cell(
                    row=idx,
                    column=col,
                ).alignment = Alignment(
                    vertical="top",
                    wrap_text=(
                        col in (
                            2,
                            5,
                        )
                    ),
                )

            svc_ws.row_dimensions[
                idx
            ].height = 24

        if ordered:
            _zebra(
                svc_ws,
                8,
                7 + len(ordered),
                1,
                6,
            )

        _autosize(
            svc_ws,
            50,
        )

        svc_ws.column_dimensions[
            "A"
        ].width = 18

        svc_ws.column_dimensions[
            "B"
        ].width = 28

        svc_ws.column_dimensions[
            "C"
        ].width = 13

        svc_ws.column_dimensions[
            "D"
        ].width = 12

        svc_ws.column_dimensions[
            "E"
        ].width = 48

        svc_ws.column_dimensions[
            "F"
        ].width = 20

        svc_ws.column_dimensions[
            "J"
        ].width = 24

        svc_ws.freeze_panes = "A8"

        svc_ws.auto_filter.ref = (
            f"A7:F{max(7, 7 + len(ordered))}"
        )

        _finish_sheet(
            svc_ws,
            landscape=True,
            fit_to_width=1,
            print_title_rows="1:7",
        )

    # Excel recalculation settings
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return output